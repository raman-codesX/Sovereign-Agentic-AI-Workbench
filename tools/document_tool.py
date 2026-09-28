from docx import Document


def create_approval_note(content, output_path):

    document = Document()

    document.add_heading(
        "EQUIPMENT INSPECTION APPROVAL NOTE",
        level=1
    )

    document.add_paragraph(content)

    document.save(output_path)

    return output_path